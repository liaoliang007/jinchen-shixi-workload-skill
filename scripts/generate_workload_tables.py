# -*- coding: utf-8 -*-
"""
2526届毕业实习指导教师工作量统计表批量生成
用法:
    python generate_workload_tables.py 何贤江            # 只生成指定教师
    python generate_workload_tables.py                    # 批量生成全部15位教师
输出: E:/openclaw/workbuddy/documents/2526届毕业实习工作量统计表/<姓名>指导学生实习工作量统计表.docx (+.pdf)
"""
import sys, shutil, copy, os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl

SRC = r'D:/何贤江指导学生实习工作量统计表.docx'
XLSX = r'D:/2526届毕业实习指导教师工作量统计表（增列教师工号学院）.xlsx'
OUTDIR = r'E:/openclaw/workbuddy/documents/2526届毕业实习工作量统计表'
SHEET = '名单教师实习明细'


def fmt(v):
    if v is None:
        return ''
    if float(v) == int(float(v)):
        return f'{float(v):.1f}'
    return f'{float(v):g}'


def clip_text(text, max_twips, font_twips=220):
    """按显示宽度截断文本（全角=1、半角=0.5字符宽），超出部分不显示。
    fixed 列宽布局下 Word 会忽略 noWrap，故直接截断文字。"""
    if not text:
        return ''
    limit = max_twips - 216  # 扣除左右单元格边距 108×2
    total = 0
    out = []
    for ch in str(text):
        w = font_twips if ord(ch) > 0x2E7F else font_twips // 2
        if total + w > limit:
            break
        total += w
        out.append(ch)
    return ''.join(out)


def load_teachers():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[SHEET]
    teachers = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        teachers.setdefault(r[1], []).append(r)
    return teachers


def set_cell(tr, idx, text):
    """替换行内第 idx 个单元格文本，保留第一个 run 的字体格式"""
    tc = tr.findall(qn('w:tc'))[idx]
    ps = tc.findall(qn('w:p'))
    first_p = ps[0]
    runs = first_p.findall(qn('w:r'))
    if runs:
        keep = runs[0]
        for r_ in runs[1:]:
            first_p.remove(r_)
        ts = keep.findall(qn('w:t'))
        for t_ in ts[1:]:
            keep.remove(t_)
        if not ts:
            t_ = OxmlElement('w:t')
            keep.append(t_)
        ts[0].text = str(text)
        ts[0].set(qn('xml:space'), 'preserve')
    for p in ps[1:]:
        tc.remove(p)


def generate(name, stu):
    total = sum(x[10] or 0 for x in stu)
    os.makedirs(OUTDIR, exist_ok=True)
    out = os.path.join(OUTDIR, f'{name}指导学生实习工作量统计表.docx')
    shutil.copy(SRC, out)
    doc = Document(out)
    body = doc.element.body

    # ---- 布局修正 ----
    # 1) 删除首段「附件4」
    for el in list(body):
        if el.tag == qn('w:p'):
            txt = ''.join(n.text or '' for n in el.iter(qn('w:t')))
            if txt.strip().startswith('附件4'):
                body.remove(el)
                break
    # 2) 落款段落（教务处+日期）移到标题下面一行
    para_els = [el for el in list(body) if el.tag == qn('w:p')]
    title_el = None
    sign_el = None
    for el in para_els:
        txt = ''.join(n.text or '' for n in el.iter(qn('w:t')))
        if '指导学生实习工作量统计表' in txt and title_el is None:
            title_el = el
        if '成都锦城学院教务处' in txt and sign_el is None:
            sign_el = el
    if title_el is not None and sign_el is not None:
        body.remove(sign_el)
        title_el.addnext(sign_el)

    # 3) 顶部页边距缩小为原来的 2/3（1800→1200缇），整体内容上移约两行
    for sec in doc.sections:
        sec.top_margin = sec.top_margin.__class__(int(sec.top_margin.twips * 2 / 3))

    # ---- 表格数据 ----
    t = doc.tables[0]
    tbl = t._tbl
    trs = tbl.findall(qn('w:tr'))
    data_tr = trs[1]
    tpl = copy.deepcopy(data_tr)
    for tc in tpl.findall(qn('w:tc')):
        tcPr = tc.find(qn('w:tcPr'))
        if tcPr is not None:
            vm = tcPr.find(qn('w:vMerge'))
            if vm is not None:
                tcPr.remove(vm)
    for tr in trs[1:]:
        tbl.remove(tr)

    for i, s in enumerate(stu, 1):
        tr = copy.deepcopy(tpl)
        vals = [i, s[3], s[4], s[5], s[6], clip_text(s[7], 2873), s[8], fmt(s[10]), '']
        for j, v in enumerate(vals):
            set_cell(tr, j, v)
        tbl.append(tr)

    # 最后一列纵向合并，填教师合计
    new_trs = tbl.findall(qn('w:tr'))
    for k, tr in enumerate(new_trs[1:], 1):
        tc = tr.findall(qn('w:tc'))[-1]
        tcPr = tc.find(qn('w:tcPr'))
        if tcPr is None:
            tcPr = OxmlElement('w:tcPr')
            tc.insert(0, tcPr)
        vm = tcPr.find(qn('w:vMerge'))
        if vm is None:
            vm = OxmlElement('w:vMerge')
            tcPr.append(vm)
        vm.set(qn('w:val'), 'restart' if k == 1 else 'continue')
        set_cell(tr, 8, fmt(total) if k == 1 else '')

    # ---- 列宽与布局调整 ----
    # 4) 去掉表格浮动定位（tblpPr/tblOverlap），表格紧邻落款行随文排列
    tblPr_el = tbl.find(qn('w:tblPr'))
    for tag in ('w:tblpPr', 'w:tblOverlap'):
        el = tblPr_el.find(qn(tag))
        if el is not None:
            tblPr_el.remove(el)
    layout = tblPr_el.find(qn('w:tblLayout'))
    if layout is not None:
        layout.set(qn('w:type'), 'fixed')
    # 5) 学年学期列加宽至单行（1230→1763），工作量合计列缩窄1/3（1599→1066），总宽不变
    grid = tbl.find(qn('w:tblGrid'))
    gcols = grid.findall(qn('w:gridCol'))
    new_w = {1: 1763, 8: 1066}
    for idx, wv in new_w.items():
        gcols[idx].set(qn('w:w'), str(wv))
    for tr in tbl.findall(qn('w:tr')):
        tcs = tr.findall(qn('w:tc'))
        for idx, wv in new_w.items():
            if idx >= len(tcs):
                continue
            tcPr = tcs[idx].find(qn('w:tcPr'))
            if tcPr is None:
                continue
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is not None:
                tcW.set(qn('w:w'), str(wv))
                tcW.set(qn('w:type'), 'dxa')

    # 6) 实习单位列（第6列，idx=5）禁止换行：超出列宽的文字裁剪不显示
    def add_nowrap(tc):
        tcPr = tc.find(qn('w:tcPr'))
        if tcPr is None:
            tcPr = OxmlElement('w:tcPr')
            tc.insert(0, tcPr)
        if tcPr.find(qn('w:noWrap')) is not None:
            return
        nw = OxmlElement('w:noWrap')
        # noWrap 须插在 tcMar/vAlign 等之前（OOXML 顺序要求）
        anchor = None
        for tag in ('w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark'):
            anchor = tcPr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(nw)
        else:
            tcPr.append(nw)
    for tr in tbl.findall(qn('w:tr')):
        tcs = tr.findall(qn('w:tc'))
        if len(tcs) > 5:
            add_nowrap(tcs[5])

    # 7) 合计列表头固定两行「工作量合计／（学时）」：字符压缩80% + 显式换行 + 缩小单元格边距
    hdr_tc = tbl.findall(qn('w:tr'))[0].findall(qn('w:tc'))[8]
    hdr_ps = hdr_tc.findall(qn('w:p'))
    first_p = hdr_ps[0]
    for p in hdr_ps[1:]:
        hdr_tc.remove(p)
    runs = first_p.findall(qn('w:r'))
    keep = runs[0]
    for r_ in runs[1:]:
        first_p.remove(r_)
    rPr = keep.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        keep.insert(0, rPr)
    for tag in ('w:w',):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    # w:w（字符缩放）须在 rFonts 之后、sz 之前
    sz_el = rPr.find(qn('w:sz'))
    w_el = OxmlElement('w:w')
    w_el.set(qn('w:val'), '80')
    if sz_el is not None:
        sz_el.addprevious(w_el)
    else:
        rPr.append(w_el)
    ts = keep.findall(qn('w:t'))
    for t_ in ts[1:]:
        keep.remove(t_)
    ts[0].text = '工作量合计'
    br = OxmlElement('w:br')
    keep.append(br)
    t2 = OxmlElement('w:t')
    t2.text = '（学时）'
    keep.append(t2)
    # 缩小表头单元格左右边距到 30 缇，保证「工作量合计」一行放下
    hdr_tcPr = hdr_tc.find(qn('w:tcPr'))
    old_mar = hdr_tcPr.find(qn('w:tcMar'))
    if old_mar is not None:
        hdr_tcPr.remove(old_mar)
    mar = OxmlElement('w:tcMar')
    for side, val in (('w:left', '30'), ('w:right', '30')):
        el = OxmlElement(side)
        el.set(qn('w:w'), val)
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    vAlign = hdr_tcPr.find(qn('w:vAlign'))
    if vAlign is not None:
        vAlign.addprevious(mar)
    else:
        hdr_tcPr.append(mar)

    doc.save(out)

    # ---- 转 PDF（Word COM）----
    # docx 可能被编辑器/预览占用，且旧 PDF 文件名可能被预览锁定：
    # 复制 docx 到临时路径，导出到带时间戳的唯一 PDF 名，再尝试归位
    pdf = out[:-5] + '.pdf'
    import win32com.client, time
    ts = time.strftime('%H%M%S')
    tmp_docx = out[:-5] + f'.tmp{ts}.docx'
    shutil.copy(out, tmp_docx)
    app = win32com.client.Dispatch('Word.Application')
    app.Visible = False
    app.DisplayAlerts = 0
    d = app.Documents.Open(tmp_docx, False, False)  # 必须可写打开，只读方式会导致导出失败
    tmp_pdf = out[:-5] + f'.new{ts}.pdf'
    try:
        d.ExportAsFixedFormat(tmp_pdf, 17)
    except Exception:
        d.SaveAs(tmp_pdf, FileFormat=17)
    d.Close(False)
    app.Quit()
    try:
        os.remove(tmp_docx)
    except OSError:
        pass
    replaced = False
    for target in (pdf, out[:-5] + '(最新版).pdf'):
        try:
            os.replace(tmp_pdf, target)
            pdf = target
            replaced = True
            break
        except PermissionError:
            continue
    if not replaced:
        pdf = tmp_pdf
        print(f'WARN 目标PDF均被占用，新版保存为: {tmp_pdf}')
    print(f'OK {name}: 学生{len(stu)}人 合计{fmt(total)}学时 -> {pdf}')
    return out, pdf


def main():
    teachers = load_teachers()
    names = sys.argv[1:] or list(teachers)
    for n in names:
        if n not in teachers:
            print(f'SKIP {n}: 名单中无此教师')
            continue
        generate(n, teachers[n])


if __name__ == '__main__':
    main()
